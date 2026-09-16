#!/usr/bin/env python3
"""
Brotherhood info-sheet template (Word .docx) with the Dream Lab footer.

Print-first: US Letter, black on white, Arial. Page 1 = info sheet,
page 2 = proposed rules. Body text is placeholder until the content lands.

  python3 build_infosheet_template.py OUT.docx [--version 1] [--date 2026-09-16]
  python3 build_infosheet_template.py OUT.docx --md cards_infosheet_draft.md --version 10

Needs python-docx + Pillow.
"""

import argparse
import io
import re
import os

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "..", "brand", "name-variants",
                    "the_dream_laboratory_oneline_black.png")

FONT = "Arial"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x59, 0x59, 0x59)
PLACEHOLDER = RGBColor(0x80, 0x80, 0x80)

BIGMO_PHONE = "(361) 423-2253"
DOMAIN = "thedreamlaboratory.org"


# ---------------------------------------------------------------- helpers

def logo_stream(max_w=1600):
    """Trim the transparent margin and downscale the 20k-px master."""
    im = Image.open(LOGO)
    im = im.crop(im.getbbox())
    if im.width > max_w:
        im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def border(el_pr, tag, sides, sz=8, color="000000", space=None):
    """Add w:pBdr / w:tcBorders sides to a pPr or tcPr."""
    box = OxmlElement(tag)
    for side in sides:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:color"), color)
        if space is not None:
            b.set(qn("w:space"), str(space))
        box.append(b)
    el_pr.append(box)


def field(paragraph, instr, size, color=GRAY):
    """Insert a Word field (PAGE, NUMPAGES) that updates on open/print."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None),
                       (None, "1"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return run


def run(p, text, size=None, bold=False, italic=False, color=None, caps=False):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.all_caps = caps
    if color is not None:
        r.font.color.rgb = color
    return r


def placeholder(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    run(p, f"[{text}]", italic=True, color=PLACEHOLDER)
    return p


def no_cell_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "nil")
        borders.append(b)
    tbl_pr.append(borders)


def cell_margins(table, left=0, right=0):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("left", left), ("right", right)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(v))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tbl_pr.append(mar)


# ---------------------------------------------------------------- styles

def force_font(style):
    """Heading/Title styles point at theme fonts; pin them to FONT."""
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in list(rf.attrib):
        del rf.attrib[k]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)


def setup_styles(doc):
    st = doc.styles

    normal = st["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.0

    # Sheet title: bold italic caps, echoing the wordmark, heavy rule below.
    title = st["Title"]
    title.font.name = FONT
    title.font.size = Pt(21)
    title.font.bold = True
    title.font.italic = True
    title.font.all_caps = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    ppr = title.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:pBdr")):
        ppr.remove(old)

    sub = st["Subtitle"]
    sub.font.name = FONT
    sub.font.size = Pt(11)
    sub.font.italic = False
    sub.font.color.rgb = GRAY
    sub.font.all_caps = True
    sub.paragraph_format.space_after = Pt(10)
    border(sub.element.get_or_add_pPr(), "w:pBdr", ["bottom"], sz=18, space=6)

    # Section heads: bold caps with a hairline under them.
    h1 = st["Heading 1"]
    h1.font.name = FONT
    h1.font.size = Pt(12.5)
    h1.font.bold = True
    h1.font.italic = False
    h1.font.all_caps = True
    h1.font.color.rgb = BLACK
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    border(h1.element.get_or_add_pPr(), "w:pBdr", ["bottom"], sz=4, space=2)

    for name in ("Title", "Subtitle", "Heading 1", "List Bullet", "List Number"):
        force_font(st[name])

    for name in ("List Bullet", "List Number"):
        s = st[name]
        s.font.name = FONT
        s.font.size = Pt(10.5)
        s.paragraph_format.space_after = Pt(3)


# ---------------------------------------------------------------- footer

def build_footer(section, version, date):
    ft = section.footer
    ft.is_linked_to_previous = False
    p0 = ft.paragraphs[0]

    width = section.page_width - section.left_margin - section.right_margin
    table = ft.add_table(rows=1, cols=3, width=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_cell_borders(table)
    cell_margins(table)

    cols = [Inches(2.45), Inches(3.05), Inches(2.0)]
    for i, w in enumerate(cols):
        table.columns[i].width = w
        table.rows[0].cells[i].width = w

    # Heavy top rule across the whole footer, matching the subtitle rule.
    for c in table.rows[0].cells:
        border(c._tc.get_or_add_tcPr(), "w:tcBorders", ["top"], sz=18)

    left, mid, right = table.rows[0].cells

    lp = left.paragraphs[0]
    lp.paragraph_format.space_before = Pt(7)
    lp.paragraph_format.space_after = Pt(0)
    lp.add_run().add_picture(logo_stream(), width=Inches(2.3))
    lp2 = left.add_paragraph()
    lp2.paragraph_format.space_after = Pt(0)
    run(lp2, DOMAIN, size=7.5, color=GRAY)

    mp = mid.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mp.paragraph_format.space_before = Pt(6)
    mp.paragraph_format.space_after = Pt(0)
    run(mp, "Fonde Brotherhood", size=9, bold=True)
    mp2 = mid.add_paragraph()
    mp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mp2.paragraph_format.space_after = Pt(0)
    run(mp2, f"Text BigMo: {BIGMO_PHONE}", size=8.5, color=GRAY)

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_before = Pt(6)
    rp.paragraph_format.space_after = Pt(0)
    run(rp, f"DRAFT v{version} — {date}", size=8.5, bold=True)
    rp2 = right.add_paragraph()
    rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp2.paragraph_format.space_after = Pt(0)
    run(rp2, "Page ", size=8.5, color=GRAY)
    field(rp2, "PAGE", 8.5)
    run(rp2, " of ", size=8.5, color=GRAY)
    field(rp2, "NUMPAGES", 8.5)

    # python-docx leaves an empty paragraph ahead of the table; move it
    # after (Word needs a paragraph to end a footer) and shrink it.
    ft._element.remove(p0._p)
    ft._element.append(p0._p)
    p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.space_before = Pt(0)
    run(p0, "", size=2)


# ---------------------------------------------------------------- body

def help_box(doc):
    """Boxed 'How you can help' callout: one-cell table, heavy black border."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "18")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tbl_pr.append(borders)
    cell_margins(t, left=180, right=180)
    c = t.rows[0].cells[0]

    hp = c.paragraphs[0]
    hp.paragraph_format.space_before = Pt(6)
    hp.paragraph_format.space_after = Pt(4)
    run(hp, "How you can help", size=12.5, bold=True, italic=True, caps=True)

    for n in range(1, 4):
        p = c.add_paragraph(style="List Number")
        run(p, f"[Ask #{n} — what you need someone to do, and by when]",
            italic=True, color=PLACEHOLDER)
    last = c.add_paragraph()
    last.paragraph_format.space_before = Pt(4)
    last.paragraph_format.space_after = Pt(6)
    run(last, "Questions? ", bold=True)
    run(last, f"Text BigMo at {BIGMO_PHONE} or ask [name].", italic=False)


def build_body(doc):
    doc.add_paragraph("Player Cards for List Management", style="Title")
    doc.add_paragraph("Experiment #1 — Fonde Brotherhood info sheet",
                      style="Subtitle")

    doc.add_heading("What we're doing", level=1)
    placeholder(doc, "Two or three sentences: the player cards, and using them "
                     "to manage the list.")

    doc.add_heading("Why", level=1)
    placeholder(doc, "The problem with how the list works today.")
    placeholder(doc, "What we hope the cards fix.")

    doc.add_heading("Experiment #1", level=1)
    for label in ("When", "Where", "Who", "How it works"):
        p = doc.add_paragraph()
        run(p, f"{label}: ", bold=True)
        run(p, "[…]", italic=True, color=PLACEHOLDER)
    for s in ("Step one", "Step two", "Step three"):
        p = doc.add_paragraph(style="List Bullet")
        run(p, f"[{s}]", italic=True, color=PLACEHOLDER)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    help_box(doc)

    # ---- page 2
    br = doc.add_paragraph()
    br.add_run().add_break(WD_BREAK.PAGE)
    br.paragraph_format.space_after = Pt(0)

    doc.add_paragraph("Proposed Rules", style="Title")
    doc.add_paragraph("Draft for discussion — not in effect", style="Subtitle")
    placeholder(doc, "One-line intro: how these rules get adopted or changed.")

    # Hand-numbered and continuous across sections, so a rule can be
    # cited as "Rule 3" and never renumbers with the help-box list.
    n = 0
    for sect in ("Getting on the list", "Order of play", "Guests"):
        doc.add_heading(sect, level=1)
        for _ in range(2):
            n += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            run(p, f"{n}.\t", bold=True)
            run(p, "[Rule]", italic=True, color=PLACEHOLDER)

    doc.add_heading("Open questions", level=1)
    p = doc.add_paragraph(style="List Bullet")
    run(p, "[Question for the Brotherhood]", italic=True, color=PLACEHOLDER)


# ---------------------------------------------------------------- markdown

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def inline(p, text, **kw):
    """**bold** and *italic* spans; [bracketed] text renders as placeholder."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run(p, part[2:-2], bold=True, **kw)
        elif part.startswith("*"):
            run(p, part[1:-1], italic=True, **kw)
        elif part.startswith("[") and part.endswith("]"):
            run(p, part, italic=True, color=PLACEHOLDER)
        else:
            run(p, part, **kw)


def build_from_md(doc, path):
    """Render the working-draft markdown subset: # title (+ next paragraph as
    subtitle), ## heads, paragraphs, - bullets, N. rules, pagebreak div.
    Page-2 rules keep their hand numbers so they match the web draft."""
    blocks, cur = [], []
    for line in open(path, encoding="utf-8").read().splitlines():
        if not line.strip():
            if cur:
                blocks.append(cur)
                cur = []
            continue
        starts_item = re.match(r"^(- |\d+\. |#)", line) or "pagebreak" in line
        if cur and starts_item:
            blocks.append(cur)
            cur = []
        cur.append(line)
    if cur:
        blocks.append(cur)

    subtitle_next = False
    break_next = False
    page = 1
    for b in blocks:
        head = b[0]
        text = " ".join(l.rstrip("\\").strip() for l in b)
        if "pagebreak" in head:
            # break *before* the next title, not a break paragraph: a full
            # page 1 would otherwise push that empty paragraph to a blank page
            page += 1
            break_next = True
            continue
        elif head.startswith("# "):
            t = doc.add_paragraph(head[2:].strip(), style="Title")
            t.paragraph_format.page_break_before = break_next
            break_next = False
            subtitle_next = True
            continue
        elif head.startswith("## "):
            doc.add_heading(head[3:].strip(), level=1)
        elif subtitle_next:
            p = doc.add_paragraph(style="Subtitle")
            # drop the web-only "Working draft N" tag; the footer carries it
            inline(p, re.split(r"\s+·\s+\*Working draft", text)[0])
        elif head.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            inline(p, text[2:])
        elif re.match(r"^\d+\. ", head):
            num, body = text.split(" ", 1)
            if page == 1:
                p = doc.add_paragraph(style="List Number")
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.3)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
                run(p, num + "\t", bold=True)
            inline(p, body)
        else:
            p = doc.add_paragraph()
            # join soft-wrapped lines first so **spans** can cross them;
            # a trailing "\\" is a hard line break
            segs, buf = [], []
            for l in b:
                buf.append(l.rstrip("\\").strip())
                if l.endswith("\\"):
                    segs.append(" ".join(buf))
                    buf = []
            if buf:
                segs.append(" ".join(buf))
            for i, seg in enumerate(segs):
                if i:
                    p.add_run().add_break()
                inline(p, seg)
        subtitle_next = False


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--version", default="1")
    ap.add_argument("--date", default="2026-09-16")
    ap.add_argument("--md", help="build from a working-draft markdown file "
                                 "instead of the placeholder template")
    a = ap.parse_args()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.5)
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.55)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    setup_styles(doc)
    build_footer(sec, a.version, a.date)
    if a.md:
        build_from_md(doc, a.md)
    else:
        build_body(doc)

    doc.core_properties.title = "Player Cards for List Management"
    doc.core_properties.author = "The Dream Laboratory"
    doc.save(a.out)
    print(a.out)


if __name__ == "__main__":
    main()
